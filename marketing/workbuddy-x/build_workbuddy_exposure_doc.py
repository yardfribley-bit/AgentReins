from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


OUT = "/Users/jatsmith/AgentSpec/AgentGuardApp/marketing/workbuddy-x/WorkBuddy-Context-Exposure.en.docx"

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.68)
sec.bottom_margin = Inches(0.68)
sec.left_margin = Inches(0.83)
sec.right_margin = Inches(0.83)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(33, 37, 41)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.13

for name, size, color, before, after in [
    ("Title", 22, RGBColor(17, 24, 39), 0, 11),
    ("Heading 1", 11, RGBColor(17, 24, 39), 12, 4),
]:
    style = styles[name]
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = color
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

title_border = OxmlElement("w:pBdr")
for side in ("top", "bottom", "left", "right"):
    element = OxmlElement(f"w:{side}")
    element.set(qn("w:val"), "nil")
    title_border.append(element)
styles["Title"]._element.get_or_add_pPr().append(title_border)

quote = styles.add_style("Evidence Quote", WD_STYLE_TYPE.PARAGRAPH)
quote.base_style = normal
quote.font.name = "Arial"
quote.font.size = Pt(9.5)
quote.font.color.rgb = RGBColor(55, 65, 81)
quote.paragraph_format.left_indent = Inches(0.18)
quote.paragraph_format.space_before = Pt(1)
quote.paragraph_format.space_after = Pt(5)

def p(text, style=None):
    return doc.add_paragraph(text, style)

def heading(text):
    doc.add_paragraph(text, "Heading 1")

doc.add_paragraph("WorkBuddy Collected My Private History for a One-Line Task", "Title")
meta = p("AgentReins Research  |  September 2026")
meta.runs[0].font.size = Pt(9)
meta.runs[0].font.color.rgb = RGBColor(107, 114, 128)

p("I asked WorkBuddy to make a WeChat publishing skill. One sentence. Before it did the work, its first model call carried 82,433 input tokens. The saved request shows what filled that space: my personal profile, my GitHub accounts, server addresses, project history, past conversations, and three local identity files. WorkBuddy assembled and included that material without asking me whether this task needed it.")

heading("The three files were not just file names")
p("The request included the actual contents of SOUL.md, IDENTITY.md, and USER.md from ~/.workbuddy/:")
p("SOUL.md told the agent that it had access to my messages, files, calendar, and possibly my home, then instructed it: “Private things stay private. Period.”", "Evidence Quote")
p("IDENTITY.md identified the agent as WorkBuddy and described its working style and strengths: code, SSH, IP networking, GitHub, credential review, and macOS automation.", "Evidence Quote")
p("USER.md named me as smithjat, identified my macOS username as jatsmith, and named https://github.com/yardfribley-bit as my usable GitHub account. It also recorded that tajleonbennis-maker had been compromised, the path to my fleet-monitor code, my interest in a location-and-spending life-log project, my preference for local data storage, and a Mac crash investigation.", "Evidence Quote")

heading("The memory layer added my infrastructure")
p("Beyond those files, WorkBuddy inserted user and Cloud Memory. In the saved prefix I found my server IPs 43.153.166.155 and 165.154.226.119, a local device at 192.168.1.39, GitHub repository and account history, SSH and deployment context, and work on AgentReins, CoPyRIT, fleet-monitor, SkillHub, and Voyager. The same prefix held at least 21 earlier user messages, 58 assistant messages, and 72 previous tool calls. This was not context for writing a WeChat publishing skill. It was a dossier attached to an ordinary request.")

heading("I am choosing to publish this. WorkBuddy did not ask.")
p("I am naming my accounts and servers here deliberately so people can see the scope of the exposure. My choice to disclose them now does not mean I approved WorkBuddy bundling them into a model request then. That distinction is the story: the product collected and forwarded far more of my personal history than the task required, while its own SOUL.md said private things should stay private.")

heading("What the evidence proves")
p("The local WorkBuddy trace contains the request prefix and the model's usage report. The first response counted 82,433 input tokens, including 80,896 cached tokens; later calls rose to 86,335, 88,200, and 91,152. The trace saves only the first 100,002 characters of generation.toolInput, so the full request may contain more. I am not claiming this trace proves the final network operator, data retention, or that OpenAI received it. It does prove WorkBuddy put these details into the model-request context.")

p("The fix starts with showing users an exact data manifest before an agent sends a request: what is included, where it came from, and why this task needs it. Anything unrelated stays out.")

source = p("Evidence and project: https://github.com/yardfribley-bit/AgentReins")
source.runs[0].font.size = Pt(9)
source.runs[0].font.color.rgb = RGBColor(37, 99, 235)

doc.core_properties.title = "WorkBuddy Collected My Private History for a One-Line Task"
doc.core_properties.author = "AgentReins Research"
doc.save(OUT)
