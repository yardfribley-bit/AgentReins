from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parent
ASSET_DIR = ROOT / "memory-security-assets"
DOCX_PATH = ROOT / "AgentReins-AI-Agent-Memory-Security.zh-CN.docx"
X_PATH = ROOT.parent / "marketing" / "workbuddy-x" / "MEMORY-SECURITY-THREAD.zh-CN.md"
ASSET_DIR.mkdir(parents=True, exist_ok=True)
X_PATH.parent.mkdir(parents=True, exist_ok=True)

FONT = "/System/Library/Fonts/STHeiti Light.ttc"
BOLD_FONT = "/System/Library/Fonts/STHeiti Medium.ttc"
MONO = "/System/Library/Fonts/SFNSMono.ttf"


def font(size, bold=False, mono=False):
    path = MONO if mono and Path(MONO).exists() else BOLD_FONT if bold else FONT
    return ImageFont.truetype(path, size=size)


def wrapped(draw, text, fnt, width):
    lines = []
    for paragraph in text.split("\n"):
        line = ""
        for ch in paragraph:
            candidate = line + ch
            if draw.textlength(candidate, font=fnt) <= width or not line:
                line = candidate
            else:
                lines.append(line)
                line = ch
        lines.append(line)
    return lines


def text_block(draw, xy, text, fnt, fill, width, spacing=12):
    x, y = xy
    for line in wrapped(draw, text, fnt, width):
        draw.text((x, y), line, font=fnt, fill=fill)
        y += fnt.size + spacing
    return y


def canvas(title, kicker, dark=False):
    bg = "#071923" if dark else "#F7F9FB"
    im = Image.new("RGB", (1600, 900), bg)
    d = ImageDraw.Draw(im)
    fg = "#F7FAFC" if dark else "#101820"
    muted = "#9DB0BD" if dark else "#52616B"
    d.text((84, 54), kicker, font=font(24, True), fill="#16C4D9")
    d.text((84, 102), title, font=font(48, True), fill=fg)
    d.line((84, 178, 1516, 178), fill="#C8D3DA" if not dark else "#24404F", width=2)
    return im, d, fg, muted


def save(im, name):
    path = ASSET_DIR / name
    im.save(path, quality=95)
    return path


def figure_1():
    im, d, fg, muted = canvas("代码上传只是显性的那一层", "热点入口  ZCode 工作区快照争议", True)
    cards = [
        (84, "公开报告中的 ZCode", "工作区源码、文档、.git 历史\n被打包为加密快照", "313 MB / 564 次重试", "#FFB547"),
        (590, "更隐蔽的日常请求", "身份、长期记忆、项目路径、\n工具读取结果进入模型上下文", "每一次普通对话", "#16C4D9"),
        (1096, "AgentReins 要回答", "具体发送了什么？谁装入的？\n与当前任务是否相关？", "内容级证据", "#56D68B"),
    ]
    for x, head, body, metric, accent in cards:
        d.rounded_rectangle((x, 244, x+420, 716), radius=18, fill="#0D2937", outline="#315264", width=2)
        d.rectangle((x, 244, x+10, 716), fill=accent)
        d.text((x+34, 282), head, font=font(30, True), fill=fg)
        y = text_block(d, (x+34, 355), body, font(27), muted, 350, 16)
        d.text((x+34, y+44), metric, font=font(31, True), fill=accent)
    d.line((504, 480, 570, 480), fill="#78909C", width=5)
    d.polygon([(570,480),(548,467),(548,493)], fill="#78909C")
    d.line((1010, 480, 1076, 480), fill="#78909C", width=5)
    d.polygon([(1076,480),(1054,467),(1054,493)], fill="#78909C")
    d.text((84, 806), "公开数字来自开发者报告；AgentReins 实验使用独立 trace 与 Canary 证据。", font=font(21), fill=muted)
    return save(im, "01-zcode-to-memory-security.png")


def figure_2():
    im, d, fg, muted = canvas("三个文件不是三个标签，而是三类长期信息", "真实上下文  已脱敏", False)
    items = [
        (84, "SOUL.md", "Agent 人格与边界", ["主动解决，不做表演式帮助", "允许形成偏好和判断", "外部动作谨慎，私密信息保密", "这些文件构成跨会话记忆"], "#0E7490"),
        (590, "IDENTITY.md", "Agent 身份与能力画像", ["名称与角色：数字搭档", "沟通风格：结论先行、直接务实", "偏好的指令交互方式", "擅长领域：代码、网络、凭据审查"], "#8B5CF6"),
        (1096, "USER.md", "用户长期画像", ["用户名、称呼与设备环境", "GitHub 账号及安全备注", "正在推进的项目与本地路径", "沟通偏好、隐私偏好、故障事项"], "#D97706"),
    ]
    for x, name, subtitle, rows, accent in items:
        d.rounded_rectangle((x, 236, x+420, 742), radius=14, fill="#FFFFFF", outline="#D7E0E5", width=2)
        d.rectangle((x, 236, x+420, 246), fill=accent)
        d.text((x+30, 282), name, font=font(34, True, mono=True), fill=accent)
        d.text((x+30, 336), subtitle, font=font(25, True), fill=fg)
        y = 402
        for row in rows:
            d.ellipse((x+31, y+7, x+43, y+19), fill=accent)
            y = text_block(d, (x+56, y), row, font(22), fg, 325, 8) + 22
    d.text((84, 808), "关键问题不是文件是否存在，而是它们为何在一个普通任务中被自动装入模型上下文。", font=font(24, True), fill="#20313B")
    return save(im, "02-three-memory-files.png")


def figure_3():
    im, d, fg, muted = canvas("用户只输入一个句号  请求却远不止一个句号", "上下文装配原理", True)
    nodes = [
        (88, 330, 278, 540, "用户输入", "。", "#16C4D9"),
        (370, 250, 790, 620, "WorkBuddy 上下文装配器", "系统指令\nSOUL 人格规则\nIDENTITY 身份\nUSER 长期画像\n历史对话\n项目路径与结构\n工具定义与连接器状态", "#FFB547"),
        (890, 330, 1120, 540, "模型请求", "73,331\n输入 tokens", "#56D68B"),
        (1220, 330, 1510, 540, "大模型", "hy3\n生成回答", "#8B5CF6"),
    ]
    for x1,y1,x2,y2,head,body,accent in nodes:
        d.rounded_rectangle((x1,y1,x2,y2), radius=18, fill="#0D2937", outline=accent, width=3)
        d.text((x1+24,y1+25), head, font=font(27, True), fill=accent)
        text_block(d, (x1+24,y1+78), body, font(25), fg, x2-x1-48, 10)
    for x1,x2 in [(278,370),(790,890),(1120,1220)]:
        d.line((x1+12,435,x2-14,435), fill="#91A7B4", width=5)
        d.polygon([(x2-14,435),(x2-36,421),(x2-36,449)], fill="#91A7B4")
    d.text((84, 756), "AgentReins 的价值：把‘自动装配’拆成可阅读的数据清单，并标注来源、触发者、接收方与相关性。", font=font(23, True), fill=fg)
    return save(im, "03-dot-to-context.png")


def figure_4():
    im, d, fg, muted = canvas("Canary 证明代码如何进入下一轮模型上下文", "同一 trace  22:27:27 至 22:27:55", False)
    steps = [
        (110, "01", "第一次模型输入", "未出现 Canary", "#64748B"),
        (560, "02", "Read 读取本地文件", "返回 404 字符\n出现 AGENTREINS_CODE_CANARY_20260922", "#D97706"),
        (1050, "03", "第二次模型输入", "再次出现 Canary\n代码已进入模型上下文", "#0E7490"),
    ]
    for x, num, head, body, accent in steps:
        d.ellipse((x, 260, x+84, 344), fill=accent)
        d.text((x+19, 275), num, font=font(28, True), fill="#FFFFFF")
        d.rounded_rectangle((x-26, 380, x+400, 660), radius=14, fill="#FFFFFF", outline="#D7E0E5", width=2)
        d.text((x, 418), head, font=font(28, True), fill=fg)
        text_block(d, (x, 484), body, font(22, mono=True), accent, 350, 12)
    for x1,x2 in [(205,520),(655,1010)]:
        d.line((x1,302,x2,302), fill="#A6B4BD", width=5)
        d.polygon([(x2,302),(x2-22,288),(x2-22,316)], fill="#A6B4BD")
    d.text((84, 780), "能证明：代码进入待发送上下文。不能仅凭本地 trace 推断最终网络接收方；去向仍需网络同流证据。", font=font(22, True), fill="#334155")
    return save(im, "04-canary-proof-chain.png")


def figure_5():
    im, d, fg, muted = canvas("真正可用的安全审计必须回答四个问题", "AgentReins 内容级证据模型", True)
    labels = [
        (100, "内容", "原文是什么\n代码 记忆 身份 历史", "#16C4D9"),
        (470, "来源", "用户主动\nAgent 自动装配\n模型要求 工具返回", "#FFB547"),
        (840, "去向", "模型域名\n记忆服务\n对象存储 其他外联", "#56D68B"),
        (1210, "判断", "与任务是否相关\n是否最小必要\n是否需要确认", "#8B5CF6"),
    ]
    for x, head, body, accent in labels:
        d.rounded_rectangle((x, 280, x+290, 650), radius=18, fill="#0D2937", outline=accent, width=3)
        d.text((x+32, 323), head, font=font(38, True), fill=accent)
        text_block(d, (x+32, 410), body, font(24), fg, 230, 16)
    for x in (400,770,1140):
        d.line((x,465,x+60,465), fill="#78909C", width=5)
        d.polygon([(x+60,465),(x+38,451),(x+38,479)], fill="#78909C")
    d.text((84, 760), "Token 数量只是规模。用户真正需要看到的是：哪些具体内容，以什么原因，被送往哪里。", font=font(26, True), fill=fg)
    return save(im, "05-agentreins-evidence-model.png")


FIGURES = [figure_1(), figure_2(), figure_3(), figure_4(), figure_5()]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0E7490")
    rpr.append(color)
    run.append(rpr)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.45
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def add_figure(doc, path, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(6.75))
    c = doc.add_paragraph(caption)
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(12)
    for run in c.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(82, 97, 107)


def build_docx():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    styles = doc.styles
    styles["Normal"].font.name = "Arial Unicode MS"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Arial Unicode MS")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial Unicode MS")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    styles["Normal"].font.size = Pt(11)
    for style_name, size in [("Title", 28), ("Heading 1", 20), ("Heading 2", 15)]:
        st = styles[style_name]
        st.font.name = "Arial Unicode MS"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Arial Unicode MS")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial Unicode MS")
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.font.bold = True

    title = doc.add_paragraph(style="Title")
    title.add_run("ZCode 代码上传风波之后")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    subtitle = doc.add_paragraph()
    subtitle.add_run("我们用 AgentReins 看见了 AI Agent 更隐蔽的数据出口").bold = True
    subtitle.runs[0].font.size = Pt(17)
    meta = doc.add_paragraph("AgentReins 实验室  2026 年 9 月 23 日  公众号发布稿")
    meta.runs[0].font.color.rgb = RGBColor(82, 97, 107)
    add_figure(doc, FIGURES[0], "图 1  ZCode 事件把代码上传带入公众视野，AgentReins 继续追问日常模型请求中的隐性数据")

    add_body(doc, "最近，ZCode 因工作区快照上传问题引发开发者关注。公开 issue 中，有开发者报告客户端将源码、文档和完整 .git 历史纳入加密快照；另一份公开分析记录了约 313 MB 的归档和 564 次失败重试。随后，厂商公开致歉、修复并开放 ZCode 源码。")
    add_body(doc, "这场争议暴露了一个长期被忽视的问题：AI 编程工具看到的远不止输入框里的那句话。整仓快照很醒目，但更日常、更难察觉的数据流，发生在每一次模型请求的上下文装配阶段。")
    add_body(doc, "我们使用 AgentReins 观察 WorkBuddy 的一次真实任务。用户要求它读取一个带有 Canary 标记的代码文件。结果不仅显示代码进入后续模型上下文，还显示三个长期记忆文件被自动装入同一请求：SOUL.md、IDENTITY.md 和 USER.md。")

    doc.add_heading("三个文件里究竟是什么", level=1)
    add_figure(doc, FIGURES[1], "图 2  三个记忆文件的真实信息类型，图中已对用户身份与路径脱敏")
    add_body(doc, "SOUL.md 不是一句抽象的“人格设定”。实测内容要求 Agent 主动解决问题、形成自己的判断、谨慎执行邮件和公开发布等外部动作，并把这些文件视为跨会话延续自身的记忆。它决定 Agent 如何行动，也定义隐私边界。")
    add_body(doc, "IDENTITY.md 进一步记录 Agent 的名字、角色、沟通风格、用户习惯的指令方式，以及代码、网络、SSH、系统服务和凭据审查等能力标签。它让每次新会话都可以迅速恢复到一个稳定的工作角色。")
    add_body(doc, "USER.md 最敏感。我们的样本包含用户称呼、macOS 用户名、设备和开发环境、GitHub 账号安全备注、多个正在推进的项目、本地工作目录、沟通偏好、隐私偏好以及仍在排查的设备故障。公开稿对具体身份、仓库和路径进行了脱敏，但 AgentReins 在本机证据中能够显示原文。")
    add_body(doc, "这些信息分别有合理用途。问题在于：当用户只提出一个与身份和历史项目无关的任务时，哪些内容仍被自动带入？用户是否知道？是否可以逐项关闭？")

    doc.add_heading("一个句号背后可能是一整个用户画像", level=1)
    add_figure(doc, FIGURES[2], "图 3  用户输入与模型请求不是同一个数据对象，中间还存在 Agent 自动装配层")
    add_body(doc, "在界面上，用户看到的是一句话，甚至只是一个句号。模型收到的却可能是一个多层消息数组：系统行为规则、身份文件、长期用户画像、项目结构、历史消息、工具定义和连接器状态，再加上用户刚刚输入的内容。")
    add_body(doc, "最新一次实测 trace 显示，模型 hy3 的输入为 73,331 tokens，其中 62,848 tokens 被标记为缓存内容，输出只有 206 tokens。数字不是文章的重点，它只说明了一件事：短输入与实际上下文规模之间可能相差几个数量级。")
    add_body(doc, "真正需要审计的是内容。SOUL.md 告诉模型该怎样扮演 Agent；IDENTITY.md 告诉模型它是谁、擅长什么；USER.md 告诉模型用户是谁、正在做什么、喜欢怎样协作。它们不是三个标签，而是完整原文。")

    doc.add_heading("代码是怎样进入下一轮请求的", level=1)
    add_figure(doc, FIGURES[3], "图 4  Canary 在同一 trace 内从本地文件移动到下一轮模型输入")
    add_body(doc, "我们在测试代码中放入唯一字符串 AGENTREINS_CODE_CANARY_20260922。第一次模型输入没有它；随后 WorkBuddy 调用 Read，工具结果中出现该字符串；第二次模型输入再次出现同一字符串。")
    add_body(doc, "这条链路证明：本地文件内容不是只在工具内部短暂存在，而是被回填进后续模型上下文。它也解释了为什么审计不能只看“Agent 调用了 Read”：安全人员还需要看到读取了什么、结果进入了哪一次请求、最终发往哪里。")
    add_body(doc, "这里必须保留证据边界。当前本地 trace 能证明代码进入待发送的模型上下文，但 trace 对超长输入保留约 100,002 字符后出现截断。接收域名和远端 IP 仍需要与同一请求的网络证据关联。AgentReins 的产品目标正是把这两条证据链合并，而不是用 Token 数替代内容。")

    doc.add_heading("ZCode 与 WorkBuddy 暴露的是同一个安全盲区", level=1)
    table = doc.add_table(rows=1, cols=4)
    table.autofit = False
    widths = [1.15, 1.95, 1.8, 1.8]
    for i, w in enumerate(widths):
        table.columns[i].width = Inches(w)
    headers = ["观察对象", "公开或实测内容", "触发方式", "用户难以回答的问题"]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        set_cell_shading(cell, "173B4C")
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255,255,255)
            run.font.bold = True
    rows = [
        ["ZCode 快照", "源码、文档、.git 历史", "公开报告称在提示前后自动生成并上传", "范围多大、能否关闭、保存多久"],
        ["WorkBuddy 记忆", "人格、Agent 身份、用户画像、项目资料", "上下文装配时自动注入", "为何与本次任务相关、是否每次携带"],
        ["工具读取结果", "本地代码和文件正文", "模型调用工具后回填下一轮", "读取结果发给了哪个模型或服务"],
    ]
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            if ri % 2:
                set_cell_shading(cells[i], "F1F6F8")
    add_body(doc, "两类事件的技术路径不同，但用户面对的是同一个问题：产品界面没有把真实数据边界讲清楚。一个“发送”按钮可能触发整仓快照，也可能触发长期记忆、历史对话和工具结果的自动装配。")

    doc.add_heading("AgentReins 为什么必须展示原文", level=1)
    add_figure(doc, FIGURES[4], "图 5  内容级审计的四个必要字段")
    add_body(doc, "“上传 73,331 tokens”只能说明规模，不能帮助用户判断风险。安全人员真正需要看到的是：其中哪一段是用户当前指令，哪一段是系统模板，哪一段来自长期记忆，哪一段是本地代码；是谁把它装入请求；它发往哪个服务；它与当前任务是否相关。")
    add_body(doc, "因此，AgentReins 的数据出境清单不应只给出数字和风险标签。它要把长上下文还原为可展开的内容清单，保留必要原文，并对身份、凭据和项目机密做本地识别与脱敏展示。用户可以先看结论，再沿证据链下钻到具体片段。")
    add_body(doc, "ZCode 风波提醒开发者关注“整个仓库是否被上传”。下一步，我们更应该关注那些每天发生、体积更小却语义更敏感的数据流：长期用户画像、历史项目、沟通习惯和由工具读取的局部代码。看见内容，才有可能真正控制内容。")

    doc.add_heading("证据与来源", level=1)
    sources = [
        ("ZCode 插件仓库 Issue 16", "https://github.com/zai-org/zcode-plugins/issues/16"),
        ("ZCode Feedback Issue 715", "https://github.com/zai-org/feedback/issues/715"),
        ("ZCode Feedback Issue 707", "https://github.com/zai-org/feedback/issues/707"),
        ("ZCode 开源仓库", "https://github.com/zai-org/ZCode"),
        ("财新关于 ZCode 开源及争议的报道", "https://www.caixinglobal.com/2026-09-22/zai-open-sources-coding-tool-after-backlash-over-silent-data-uploads-102487378.html"),
    ]
    for label, url in sources:
        p = doc.add_paragraph(style=None)
        p.add_run("• ")
        add_hyperlink(p, label, url)
    add_body(doc, "本文涉及 WorkBuddy 的内容来自 2026 年 9 月 22 日 AgentReins 本机授权实验。公开图片已脱敏；原始 trace、时间和 Canary 证据保留在实验设备中。本文不把本地上下文记录误称为网络接收证明。")
    doc.save(DOCX_PATH)


def build_x():
    text = """# X 线程发布稿

## 主帖

ZCode 的代码上传风波，让大家第一次认真追问：AI 编程工具到底从电脑里拿走了什么？

我们用 AgentReins 检查了另一个更隐蔽的数据出口：模型请求里的长期记忆。

用户只发一句话，Agent 可能同时带上人格、身份、用户画像、项目资料和本地代码。🧵 1/7

配图：`docs/memory-security-assets/01-zcode-to-memory-security.png`

## 2/7

公开 issue 报告称，ZCode 将工作区源码、文档和 `.git` 历史纳入加密快照；一个样本约 313 MB，并记录了 564 次重试。

这是整仓上传。我们继续追问：没有整仓打包时，Agent 每次请求又在发送什么？

来源：GitHub Issue #16、#707、#715。

## 3/7

AgentReins 在 WorkBuddy 的真实上下文里看到三个文件：

SOUL.md：Agent 的人格、行动原则和隐私边界  
IDENTITY.md：Agent 身份、能力和沟通习惯  
USER.md：用户称呼、设备、账号备注、活跃项目、本地路径和长期偏好

这不是三个标签，是完整原文。

配图：`docs/memory-security-assets/02-three-memory-files.png`

## 4/7

用户只输入一个句号，模型看到的也可能是：

系统指令 + SOUL + IDENTITY + USER + 历史对话 + 项目结构 + 工具定义 + 连接器状态。

最新样本输入 73,331 tokens，输出只有 206 tokens。数字只是规模，内容才是风险。

配图：`docs/memory-security-assets/03-dot-to-context.png`

## 5/7

我们又做了 Canary 实验：

① 第一次模型输入没有标记  
② WorkBuddy 调用 Read，读取带标记的本地代码  
③ 第二次模型输入出现同一标记

这证明本地代码进入了下一轮模型上下文。

配图：`docs/memory-security-assets/04-canary-proof-chain.png`

## 6/7

但我们不会把本地 trace 冒充网络接收证明。

当前证据证明“代码进入待发送上下文”；接收域名和远端 IP 仍需同一请求的网络证据关联。超长 trace 还存在约 100,002 字符的截断。

安全研究最重要的是：结论必须和证据强度一致。

## 7/7

AgentReins 不只告诉你“上传了多少 tokens”。

它要回答四件事：
内容是什么、来源是什么、去了哪里、是否有必要。

AI Agent 的数据边界不应藏在一次请求的几十万字符里。用户应该看得见，也应该能控制。

配图：`docs/memory-security-assets/05-agentreins-evidence-model.png`

## 发布配置

- 首帖使用图 1，避免一上来堆技术细节。
- 第 3、4、5、7 条分别使用图 2 至图 5。
- 不在图片或正文中公开真实用户名、绝对路径、仓库名、Token 或凭据。
- ZCode 数字始终写成“公开 issue 报告”或“开发者样本”，不要写成 AgentReins 已独立复现。
- WorkBuddy 结论写成“进入模型上下文”；获得网络同流证据前，不写“服务器已经接收”。
- 建议发布时间：工作日 11:30 或 20:30；首帖 24 小时内置顶。
"""
    X_PATH.write_text(text, encoding="utf-8")


build_docx()
build_x()
print(DOCX_PATH)
print(X_PATH)
