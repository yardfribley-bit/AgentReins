from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "ZCODE-34-TOOLS-NETWORK-DECODE.zh-CN.md"
OUTPUT = ROOT / "AgentReins-ZCode-34-Tools-Network-Decode.zh-CN.docx"

NAVY = "17365D"
BLUE = "2F6B9A"
PALE = "EEF4F8"
LIGHT = "F7F9FB"
GRID = "D9E1E8"
TEXT = RGBColor(37, 43, 51)
MUTED = RGBColor(91, 101, 113)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    relation = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.extend([color, underline])
    run.append(props)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_inline(paragraph, text, bold_all=False):
    pattern = re.compile(r"(https?://\S+|`[^`]+`|\*\*[^*]+\*\*)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            run.bold = bold_all
        token = match.group(0)
        if token.startswith("http"):
            add_hyperlink(paragraph, token.rstrip("。；，,)"), token.rstrip("。；，,)"))
            suffix = token[len(token.rstrip("。；，,)")):]
            if suffix:
                paragraph.add_run(suffix)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Menlo"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Menlo")
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(36, 64, 89)
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "EEF1F4")
            run._element.get_or_add_rPr().append(shd)
        else:
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        run.bold = bold_all


def set_keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial Unicode MS"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    normal.font.size = Pt(11)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.line_spacing = 1.45
    normal.paragraph_format.space_after = Pt(7)

    title = styles["Title"]
    title.font.name = "Arial Unicode MS"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(12)

    for name, size, before, after in (("Heading 1", 17, 22, 8), ("Heading 2", 13.5, 16, 6)):
        style = styles[name]
        style.font.name = "Arial Unicode MS"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def force_cjk_fonts(doc):
    def apply(paragraph):
        for run in paragraph.runs:
            run.font.name = "Arial Unicode MS"
            fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
            fonts.set(qn("w:ascii"), "Arial Unicode MS")
            fonts.set(qn("w:hAnsi"), "Arial Unicode MS")
            fonts.set(qn("w:eastAsia"), "Arial Unicode MS")
            fonts.set(qn("w:cs"), "Arial Unicode MS")
    for paragraph in doc.paragraphs:
        apply(paragraph)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    apply(paragraph)


def parse_table(lines, start):
    rows = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        cells = [item.strip() for item in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            rows.append(cells)
        index += 1
    return rows, index


def add_table(doc, rows):
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_shading(cell, NAVY if r_idx == 0 else (PALE if r_idx % 2 == 0 else "FFFFFF"))
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.2
            if c_idx > 0 and len(value) < 22:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(paragraph, value, bold_all=(r_idx == 0))
            for run in paragraph.runs:
                run.font.size = Pt(9.5 if len(rows) > 8 else 10)
                if r_idx == 0:
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.bold = True
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    configure_styles(doc)

    index = 0
    paragraph_count = 0
    while index < len(lines):
        raw = lines[index]
        text = raw.strip()
        if not text:
            index += 1
            continue
        if text.startswith("|" ) and index + 1 < len(lines) and lines[index + 1].strip().startswith("|"):
            rows, index = parse_table(lines, index)
            add_table(doc, rows)
            continue
        if text.startswith("# "):
            p = doc.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_inline(p, text[2:])
        elif text.startswith("## "):
            p = doc.add_paragraph(style="Heading 1")
            add_inline(p, text[3:])
        elif text.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            add_inline(p, text[4:])
        elif text.startswith("**副标题："):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(7)
            run = p.add_run(text.replace("**", ""))
            run.bold = True
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(47, 107, 154)
        elif text.startswith("作者："):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(16)
            run = p.add_run(text)
            run.font.size = Pt(9.5)
            run.font.color.rgb = MUTED
        elif re.match(r"^\d+\. ", text):
            match = re.match(r"^(\d+)\. (.*)$", text)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.18)
            p.paragraph_format.first_line_indent = Inches(-0.18)
            p.add_run(f"{match.group(1)}. ").bold = True
            add_inline(p, match.group(2))
        elif text.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, text[2:])
        elif text.startswith("`") and text.endswith("`") and text.count("`") == 2:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.28)
            p.paragraph_format.right_indent = Inches(0.28)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(text[1:-1])
            run.font.name = "Menlo"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Menlo")
            run.font.size = Pt(9.5)
            set_cell = OxmlElement("w:shd")
            set_cell.set(qn("w:fill"), "EEF4F8")
            p._p.get_or_add_pPr().append(set_cell)
        else:
            p = doc.add_paragraph()
            add_inline(p, text)
            paragraph_count += 1
            if paragraph_count == 1:
                for run in p.runs:
                    run.bold = True
                    run.font.size = Pt(11.5)
                p.paragraph_format.space_after = Pt(12)
        index += 1

    core = doc.core_properties
    core.title = "通过ZCode网络解码分析出34个ZCode工具"
    core.subject = "ZCode 工具 Schema、执行链与 Agent 最小权限研究"
    core.author = "AgentReins 研究组"
    core.keywords = "ZCode, AgentReins, AI Agent, 工具调用, 上下文安全, Model I/O"
    force_cjk_fonts(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
